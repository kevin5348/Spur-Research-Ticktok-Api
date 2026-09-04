
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - only runs when dependency is absent
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with: pip install openpyxl"
    ) from exc


DEFAULT_IRRELEVANT_FLAGS = {"irrelevant", "irrelvant"}


def normalise_text(value: Any) -> str:
    """Normalise a label for case-insensitive comparison."""
    return str(value if value is not None else "").strip().casefold()


def normalise_id(value: Any) -> str:
    """Convert an Excel/JSON video ID to a comparable string."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).casefold()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def find_column(headers: list[Any], requested: str) -> int:
    """Return the zero-based index of a column, ignoring case and whitespace."""
    wanted = normalise_text(requested)
    for index, header in enumerate(headers):
        if normalise_text(header) == wanted:
            return index
    available = ", ".join(str(h) for h in headers if h is not None)
    raise ValueError(
        f"Column {requested!r} was not found. Available columns: {available}"
    )


def read_irrelevant_ids(
    excel_path: Path,
    sheet_name: str | None,
    id_column: str,
    flag_column: str,
    irrelevant_flags: set[str],
) -> tuple[set[str], int, str]:
    """Read unique video IDs whose flag matches an irrelevant flag."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet {sheet_name!r} was not found. "
                    f"Available sheets: {', '.join(workbook.sheetnames)}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook[workbook.sheetnames[0]]

        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration as exc:
            raise ValueError("The selected Excel sheet is empty.") from exc

        id_index = find_column(headers, id_column)
        flag_index = find_column(headers, flag_column)
        accepted = {normalise_text(flag) for flag in irrelevant_flags}

        irrelevant_ids: set[str] = set()
        matched_rows = 0
        missing_ids = 0

        for row in rows:
            flag = row[flag_index] if flag_index < len(row) else None
            if normalise_text(flag) not in accepted:
                continue
            matched_rows += 1
            value = row[id_index] if id_index < len(row) else None
            video_id = normalise_id(value)
            if video_id:
                irrelevant_ids.add(video_id)
            else:
                missing_ids += 1

        if missing_ids:
            print(
                f"Warning: {missing_ids} irrelevant Excel row(s) had no video ID.",
                file=sys.stderr,
            )

        return irrelevant_ids, matched_rows, worksheet.title
    finally:
        workbook.close()


def get_company_records(data: Any) -> list[dict[str, Any]]:
    """Return mutable company objects from one of the supported JSON layouts."""
    if isinstance(data, list):
        records = data
    elif isinstance(data, dict) and isinstance(data.get("companies"), list):
        records = data["companies"]
    elif isinstance(data, dict) and isinstance(data.get("videos"), list):
        records = [data]
    else:
        raise ValueError(
            "Unsupported JSON structure. Expected one company object, a list of "
            "company objects, or an object with a 'companies' list."
        )

    for index, record in enumerate(records):
        if not isinstance(record, dict):
            raise ValueError(f"Company entry {index} is not a JSON object.")
        if not isinstance(record.get("videos"), list):
            raise ValueError(f"Company entry {index} does not contain a 'videos' list.")
    return records


def comment_count(video: dict[str, Any]) -> int:
    """Count the comment objects stored under one video."""
    comments = video.get("comments", [])
    if comments is None:
        return 0
    if not isinstance(comments, list):
        video_id = video.get("video_id", "<missing>")
        raise ValueError(f"Video {video_id} has a 'comments' value that is not a list.")
    return len(comments)


def filter_companies(
    records: Iterable[dict[str, Any]], irrelevant_ids: set[str]
) -> tuple[dict[str, dict[str, int]], set[str]]:
    """Filter videos in place and return aggregated before/after counts."""
    stats: dict[str, dict[str, int]] = defaultdict(
        lambda: {
            "videos_before": 0,
            "videos_removed": 0,
            "videos_after": 0,
            "comments_before": 0,
            "comments_removed": 0,
            "comments_after": 0,
        }
    )
    removed_ids: set[str] = set()

    for record_number, record in enumerate(records, start=1):
        company = str(record.get("company") or f"<company {record_number}>")
        videos = record["videos"]
        kept_videos: list[dict[str, Any]] = []

        for video_number, video in enumerate(videos, start=1):
            if not isinstance(video, dict):
                raise ValueError(
                    f"Video entry {video_number} for {company!r} is not a JSON object."
                )

            comments = comment_count(video)
            stats[company]["videos_before"] += 1
            stats[company]["comments_before"] += comments

            video_id = normalise_id(video.get("video_id"))
            if video_id and video_id in irrelevant_ids:
                stats[company]["videos_removed"] += 1
                stats[company]["comments_removed"] += comments
                removed_ids.add(video_id)
            else:
                kept_videos.append(video)
                stats[company]["videos_after"] += 1
                stats[company]["comments_after"] += comments

        record["videos"] = kept_videos
        record["video_count"] = len(kept_videos)

    return dict(stats), removed_ids


def print_count_table(stats: dict[str, dict[str, int]]) -> None:
    """Print per-company and total video/comment counts."""
    columns = [
        ("Company", "company"),
        ("Videos before", "videos_before"),
        ("Removed", "videos_removed"),
        ("Videos after", "videos_after"),
        ("Comments before", "comments_before"),
        ("Removed comments", "comments_removed"),
        ("Comments after", "comments_after"),
    ]

    rows: list[dict[str, Any]] = []
    for company in sorted(stats, key=str.casefold):
        rows.append({"company": company, **stats[company]})

    total = {key: 0 for _, key in columns if key != "company"}
    for values in stats.values():
        for key in total:
            total[key] += values[key]
    rows.append({"company": "TOTAL", **total})

    def display_value(row: dict[str, Any], key: str) -> str:
        value = row[key]
        return str(value) if key == "company" else f"{value:,}"

    widths: dict[str, int] = {}
    for title, key in columns:
        widths[key] = max(len(title), *(len(display_value(row, key)) for row in rows))

    header = "  ".join(title.ljust(widths[key]) for title, key in columns)
    divider = "  ".join("-" * widths[key] for _, key in columns)
    print("\n" + header)
    print(divider)
    for row in rows:
        cells = []
        for _, key in columns:
            value = display_value(row, key)
            if key == "company":
                cells.append(value.ljust(widths[key]))
            else:
                cells.append(value.rjust(widths[key]))
        print("  ".join(cells))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Remove videos from JSON when their video ID is flagged irrelevant "
            "in an Excel workbook."
        )
    )
    parser.add_argument("excel", type=Path, help="Excel file containing IDs and flags")
    parser.add_argument("input_json", type=Path, help="Original company/video JSON file")
    parser.add_argument("output_json", type=Path, help="Filtered JSON output file")
    parser.add_argument(
        "--sheet",
        help="Excel sheet name (default: the first sheet)",
    )
    parser.add_argument(
        "--id-column",
        default="id",
        help="Excel video ID column (default: id)",
    )
    parser.add_argument(
        "--flag-column",
        default="v3_flag",
        help="Excel flag column (default: v3_flag)",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()

    for path, label in ((args.excel, "Excel"), (args.input_json, "JSON")):
        if not path.is_file():
            raise FileNotFoundError(f"{label} input file was not found: {path}")
    if args.input_json.resolve() == args.output_json.resolve():
        raise ValueError(
            "The output JSON must have a different path so the original is preserved."
        )

    irrelevant_ids, matched_rows, sheet_used = read_irrelevant_ids(
        excel_path=args.excel,
        sheet_name=args.sheet,
        id_column=args.id_column,
        flag_column=args.flag_column,
        irrelevant_flags=DEFAULT_IRRELEVANT_FLAGS,
    )

    with args.input_json.open("r", encoding="utf-8-sig") as file:
        data = json.load(file)

    records = get_company_records(data)
    stats, removed_ids = filter_companies(records, irrelevant_ids)

    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    with args.output_json.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")

    print(f"Excel sheet: {sheet_used}")
    print(f"Excel rows marked irrelevant: {matched_rows:,}")
    print(f"Unique irrelevant IDs in Excel: {len(irrelevant_ids):,}")
    print(f"Unique matching IDs removed from JSON: {len(removed_ids):,}")
    print(f"Irrelevant Excel IDs not present in JSON: {len(irrelevant_ids - removed_ids):,}")
    print_count_table(stats)
    print(f"\nFiltered JSON written to: {args.output_json}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError, json.JSONDecodeError) as error:
        print(f"Error: {error}", file=sys.stderr)
        raise SystemExit(1)